from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
source_md = ROOT / 'NQMS_DOCUMENTATION.md'
output_docx = ROOT / 'NQMS Documentation.docx'

award_guidance = [
    'To successfully secure the Schools Division Superintendent (SDS) endorsement from the DepEd Bislig City Division for the National Quality Management System (NQMS) web app, follow formal administrative routing and presentation protocols required for the Gawad Regional Director data transformation category.',
    'Because this is an entry for an official regional award, securing the SDS signature requires demonstrating compliance, system security, and data privacy standardizations.',
    '',
    '1. Technical Preparation & Compliance Check',
    '• Data Privacy and Governance Compliance: Clearly explain compliance with the Data Privacy Act of 2012 and include a Data Protection Officer (DPO) statement if personnel data is processed.',
    '• Hosting and Accessibility: Ensure the web app is deployed on staging/production and provide secure temporary reviewer credentials in the documentation for SDS technical validation.',
    '• Data Transformation Narrative: Include before-and-after metrics such as processing time reduction, paper backlog elimination, and data accuracy improvements.',
    '',
    '2. Administrative Routing Protocol',
    '1) Division Information Technology Officer (ITO): Conducts technical assessment and issues technical endorsement.',
    '2) Quality Management Representative (QMR) / Chief of Functional Division: Validates alignment with NQMS and division SOP implementation.',
    '3) Assistant Schools Division Superintendent (ASDS): Performs final administrative completeness check before elevation to SDS.',
    '',
    '3. Securing the SDS Signatories',
    '• Include signatory blocks in the Project Proposal, Technical Documentation, and Award Entry Form.',
    '• Submission packet should contain entry form, full system documentation, endorsement letters from ITO/QMR, and a one-page executive summary highlighting data transformation gains.',
]


def add_centered_title(document: Document, text: str, size: int = 20) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(size)


def add_signatory_page(document: Document) -> None:
    add_centered_title(document, 'NQMS Documentation', 24)

    subtitle = document.add_paragraph('Signatories')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].bold = True
    subtitle.runs[0].font.size = Pt(14)

    document.add_paragraph('')
    document.add_paragraph('RECOMMENDING APPROVAL:').runs[0].bold = True
    document.add_paragraph('(Name of ITO / QMR Chief)')
    document.add_paragraph('(Position Title)')

    document.add_paragraph('')
    document.add_paragraph('APPROVED:').runs[0].bold = True
    document.add_paragraph('[INSERT CURRENT SDS NAME]')
    document.add_paragraph('Schools Division Superintendent')
    document.add_paragraph('Schools Division of Bislig City')

    document.add_page_break()


def add_award_guidance(document: Document) -> None:
    heading = document.add_paragraph('SDS Endorsement and Administrative Routing Guide')
    heading.runs[0].bold = True
    heading.runs[0].font.size = Pt(14)

    for line in award_guidance:
        if not line:
            document.add_paragraph('')
            continue

        if line.startswith('1. ') or line.startswith('2. ') or line.startswith('3. '):
            p = document.add_paragraph(line)
            p.runs[0].bold = True
        else:
            document.add_paragraph(line)

    document.add_page_break()


def add_markdown_content(document: Document, markdown_text: str) -> None:
    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()
        if not line:
            document.add_paragraph('')
            continue

        if line.startswith('### '):
            p = document.add_paragraph(line[4:])
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(13)
            continue

        if line.startswith('## '):
            p = document.add_paragraph(line[3:])
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(15)
            continue

        if line.startswith('# '):
            p = document.add_paragraph(line[2:])
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(17)
            continue

        if line.startswith('- '):
            document.add_paragraph('• ' + line[2:])
            continue

        if line.startswith('```'):
            continue

        document.add_paragraph(line)


def main() -> None:
    if not source_md.exists():
        raise FileNotFoundError(f'Missing source file: {source_md}')

    markdown_text = source_md.read_text(encoding='utf-8')

    doc = Document()
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)

    add_signatory_page(doc)
    add_award_guidance(doc)
    add_markdown_content(doc, markdown_text)

    doc.save(output_docx)
    print(f'Created: {output_docx}')


if __name__ == '__main__':
    main()
